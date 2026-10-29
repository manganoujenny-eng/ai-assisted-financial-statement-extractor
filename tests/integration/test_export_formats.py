"""The analysis export, as a PDF report or as an Excel workbook.

GET /api/cases/<id>/export?format=pdf|xlsx&lang=en|fr

Both open with the same summary (what the analysis says: strengths,
weaknesses, what the figures can be relied on for, points of attention —
see tests/unit/test_report_content.py), both are in the language of the
screen (``lang``), and both carry the AI commentary ONLY once an analyst
approved it (BR-13).

* xlsx (the default, as before): sheets Summary, Ratios, Checks,
  Validated data (+ AI commentary when approved);
* pdf: the same, as a report to hand over.
"""

from __future__ import annotations

import io

import fitz  # PyMuPDF — already a dependency (document_processing)
import pytest
from openpyxl import load_workbook

from app import create_app
from app.config import Config

from integration.test_api_end_to_end import FakeReader, FakeStructurer
from integration.test_intergration_api import FakeWriter


@pytest.fixture
def app(tmp_path, balanced_amounts):
    app = create_app(Config(storage="memory", upload_dir=tmp_path))
    c = app.config["container"]
    c.extraction.readers = [FakeReader()]
    c.extraction.structurer = FakeStructurer(balanced_amounts)
    c.insight.writer = FakeWriter()
    return app


def _analysed_case(http):
    case_id = http.post("/api/cases", json={"client": "SOCATRANS", "fiscal_year": 2024,
                                            "analyst": "Helena"}).get_json()["id"]
    http.post(f"/api/cases/{case_id}/documents",
              data={"file": (io.BytesIO(b"%PDF-1.7 export"), "statement.pdf")},
              content_type="multipart/form-data")
    run_id = http.post(f"/api/cases/{case_id}/extractions").get_json()["id"]
    http.post(f"/api/cases/{case_id}/submission", json={})
    dataset_id = http.post(f"/api/extractions/{run_id}/validation",
                           json={"supervisor": "Paul", "acknowledge_failed_checks": True}).get_json()["id"]
    analysed = http.post(f"/api/datasets/{dataset_id}/analysis").get_json()
    return case_id, analysed["ai_analysis"]


def _text(pdf_bytes: bytes) -> str:
    with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
        text = "\n".join(page.get_text() for page in doc)
    # The PDF font draws "fi" / "fl" as one ligature glyph.
    return text.replace("ﬁ", "fi").replace("ﬂ", "fl")


def test_excel_stays_the_default(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    response = http.get(f"/api/cases/{case_id}/export")
    assert response.status_code == 200
    assert response.mimetype.endswith("spreadsheetml.sheet")
    assert response.data[:2] == b"PK"  # a zip: the .xlsx container


def test_a_pdf_report_in_english(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    response = http.get(f"/api/cases/{case_id}/export?format=pdf&lang=en")
    assert response.status_code == 200 and response.mimetype == "application/pdf"
    assert response.data[:4] == b"%PDF"
    assert "SOCATRANS_2024_analysis.pdf" in response.headers["Content-Disposition"]
    text = _text(response.data)
    for expected in ("SOCATRANS", "2024", "Current ratio", "Return on equity",
                     "Balance sheet equilibrium", "TOTAL ACTIF CIRCULANT", "Page 1"):
        assert expected in text, expected


def test_a_pdf_report_in_french(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    text = _text(http.get(f"/api/cases/{case_id}/export?format=pdf&lang=fr").data)
    for expected in ("Ratio de liquidité générale", "Équilibre du bilan", "Contrôles de cohérence"):
        assert expected in text, expected


def test_the_ai_commentary_is_in_the_report_only_once_approved(app):
    http = app.test_client()
    case_id, draft = _analysed_case(http)
    before = _text(http.get(f"/api/cases/{case_id}/export?format=pdf").data)
    assert "Draft overview." not in before
    http.post(f"/api/ai-analyses/{draft['id']}/approval",
              json={"author": "Helena", "sections": {**draft["sections"], "overview": "Approved overview."}})
    after = _text(http.get(f"/api/cases/{case_id}/export?format=pdf").data)
    assert "Approved overview." in after and "Helena" in after


def _workbook(http, case_id, lang=None):
    url = f"/api/cases/{case_id}/export?format=xlsx" + (f"&lang={lang}" if lang else "")
    return load_workbook(io.BytesIO(http.get(url).data))


def _cells(sheet) -> str:
    return "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row
                     if cell.value is not None)


def test_the_pdf_opens_with_the_summary(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    text = _text(http.get(f"/api/cases/{case_id}/export?format=pdf&lang=en").data)
    for heading in ("Summary", "Strengths", "Weaknesses", "What the figures can be relied on for",
                    "Points of attention", "Ratios in detail"):
        assert heading in text, heading
    assert text.index("Summary") < text.index("Ratios in detail")
    french = _text(http.get(f"/api/cases/{case_id}/export?format=pdf&lang=fr").data)
    for heading in ("Synthèse", "Points forts", "Points faibles", "Points d'attention"):
        assert heading in french, heading


def test_the_workbook_in_english(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    book = _workbook(http, case_id, "en")
    assert book.sheetnames == ["Summary", "Ratios", "Checks", "Validated data"]
    summary = _cells(book["Summary"])
    for expected in ("SOCATRANS", "Strengths", "Weaknesses", "Points of attention"):
        assert expected in summary, expected
    assert "Current ratio" in _cells(book["Ratios"])
    assert "Balance sheet equilibrium" in _cells(book["Checks"])


def test_the_workbook_follows_the_screen_language(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    book = _workbook(http, case_id, "fr")
    assert book.sheetnames == ["Synthèse", "Ratios", "Contrôles", "Données validées"]
    summary = _cells(book["Synthèse"])
    for expected in ("Points forts", "Points faibles", "Points d'attention", "Exercice"):
        assert expected in summary, expected
    assert "Ratio de liquidité générale" in _cells(book["Ratios"])
    assert "Équilibre du bilan" in _cells(book["Contrôles"])
    # The amounts are numbers, so Excel can compute with them.
    amounts = [row[2].value for row in book["Données validées"].iter_rows(min_row=2)
               if isinstance(row[2].value, (int, float))]
    assert amounts


def test_the_workbook_carries_the_approved_ai_commentary(app):
    http = app.test_client()
    case_id, draft = _analysed_case(http)
    assert "AI commentary" not in _workbook(http, case_id, "en").sheetnames
    http.post(f"/api/ai-analyses/{draft['id']}/approval",
              json={"author": "Helena", "sections": {**draft["sections"], "overview": "Approved overview."}})
    book = _workbook(http, case_id, "en")
    assert book.sheetnames[-1] == "AI commentary"
    assert "Approved overview." in _cells(book["AI commentary"])


def test_an_unknown_language_falls_back_to_english(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    assert _workbook(http, case_id, "de").sheetnames[0] == "Summary"


def test_an_unknown_format_is_refused(app):
    http = app.test_client()
    case_id, _ = _analysed_case(http)
    assert http.get(f"/api/cases/{case_id}/export?format=docx").status_code == 422