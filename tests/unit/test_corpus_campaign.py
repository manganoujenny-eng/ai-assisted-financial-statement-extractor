"""Week 7 — the corpus campaign (services/corpus_campaign.py).

Every test writes a small annotation workbook in the RP001 format into a
temporary folder and uses a fake extractor, so the expected answer can be
worked out by hand and no PDF is needed. ``BALANCED`` (conftest.py) is the
hand reading: it balances, and every check passes on it.
"""

from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.services.corpus_campaign import (
    AnnotationError,
    Extraction,
    load_annotation,
    run_campaign,
)
from conftest import BALANCED


def write_annotation(path, amounts, year=2024, extra_rows=()):
    """A workbook laid out like annotation_RP001.xlsx (header on row 4)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Annotation"
    ws["B2"] = "ANNOTATIONS LIST"
    ws.append([])
    ws.append([None, "Statement", "code", " Standard Field", "Original Label", "Value ",
               "Unit", "Currency", "Year", "Page", "Status", "Mapping note"])
    for code, value in amounts.items():
        ws.append([None, "Balance Sheet ", code, "", "", value, "", "", year, 1, "Confirmed", ""])
    for row in extra_rows:
        ws.append(row)
    wb.save(path)


def package(corpus, name, amounts, **kwargs):
    folder = corpus / name
    folder.mkdir()
    (folder / "statement.pdf").write_bytes(b"%PDF-1.4 placeholder")
    write_annotation(folder / f"annotation_{name}.xlsx", amounts, **kwargs)
    return folder


def extractor(amounts_by_package):
    """A fake extractor: returns what we decide the system 'read'."""
    def extract(documents):
        name = documents[0].parent.name
        return Extraction({k: Decimal(v) for k, v in amounts_by_package[name].items()},
                          engine="FAKE")
    return extract


# --- reading the annotation workbook --------------------------------------
class TestLoadAnnotation:
    def test_reads_codes_and_values_by_their_header(self, tmp_path):
        write_annotation(tmp_path / "a.xlsx", {"BK": 100, "XI": -452948553})
        annotation = load_annotation(tmp_path / "a.xlsx")
        assert annotation.amounts == {"BK": Decimal(100), "XI": Decimal(-452948553)}
        assert annotation.year == 2024

    def test_a_text_amount_with_spaces_is_read(self, tmp_path):
        write_annotation(tmp_path / "a.xlsx", {"BZ": "7 969 151 104"})
        assert load_annotation(tmp_path / "a.xlsx").amounts["BZ"] == Decimal(7969151104)

    def test_an_unreadable_line_is_reported_not_guessed(self, tmp_path):
        write_annotation(tmp_path / "a.xlsx", {"BK": 100, "CP": "see page 2"})
        annotation = load_annotation(tmp_path / "a.xlsx")
        assert "CP" not in annotation.amounts
        assert any("CP" in p for p in annotation.problems)

    def test_a_workbook_without_the_header_is_refused(self, tmp_path):
        wb = Workbook()
        wb.active.append(["nothing", "here"])
        wb.save(tmp_path / "a.xlsx")
        with pytest.raises(AnnotationError):
            load_annotation(tmp_path / "a.xlsx")


# --- one campaign ---------------------------------------------------------
class TestCampaign:
    def test_a_perfect_extraction(self, tmp_path):
        package(tmp_path, "RP001", BALANCED)
        campaign = run_campaign(tmp_path, extractor({"RP001": BALANCED}))
        s = campaign.summary()
        assert s["evaluated"] == 1
        assert s["field_accuracy"] == 1 and s["ratio_accuracy"] == 1
        assert s["amplification"] is None
        assert s["silent_false_positives"] == []

    def test_a_wrong_hand_reading_is_rejected_not_measured(self, tmp_path):
        # The RP001 lesson: a reference that does not balance would make a
        # correct extraction look wrong. It is left out, with the reason.
        wrong = dict(BALANCED, AZ=5_699_223_032)  # gross instead of net
        package(tmp_path, "RP001", wrong)
        campaign = run_campaign(tmp_path, extractor({"RP001": BALANCED}))
        assert campaign.packages[0].status == "REFERENCE_REJECTED"
        assert "CHK002" in campaign.packages[0].reason
        assert campaign.summary()["evaluated"] == 0

    def test_a_silent_revenue_error_is_found(self, tmp_path):
        package(tmp_path, "RP001", BALANCED)
        read = dict(BALANCED, XB=22_000_000)  # +10 %, no check sees it
        p = run_campaign(tmp_path, extractor({"RP001": read})).packages[0]
        assert p.silent_false_positives == ["XB"]
        assert [w["code"] for w in p.wrong_fields] == ["XB"]
        assert [w["code"] for w in p.wrong_ratios] == ["R4"]
        # 14 items judged (BALANCED's XG is not one of them): (1/5) / (1/14)
        assert p.amplification == Decimal("2.80")

    def test_fields_and_ratios_are_pooled_over_the_corpus(self, tmp_path):
        package(tmp_path, "RP001", BALANCED)
        package(tmp_path, "RP002", BALANCED)
        read = {"RP001": BALANCED, "RP002": dict(BALANCED, XB=22_000_000)}
        s = run_campaign(tmp_path, extractor(read)).summary()
        assert (s["fields_judged"], s["fields_wrong"]) == (28, 1)
        assert (s["ratios_judged"], s["ratios_wrong"]) == (10, 1)
        assert s["amplification"] == Decimal("2.80")  # (1/10) / (1/28)
        assert s["silent_false_positives"] == [{"package": "RP002", "code": "XB"}]

    def test_lines_outside_the_judged_items_are_compared_apart(self, tmp_path):
        package(tmp_path, "RP001", dict(BALANCED, BA=10, XG=1_500_000))
        read = dict(BALANCED, BA=10, XG=1_400_000)
        p = run_campaign(tmp_path, extractor({"RP001": read})).packages[0]
        assert p.field_accuracy == 1  # BA, XG are not in the measurements
        assert p.other_lines["annotated"] == 2
        assert [d["code"] for d in p.other_lines["different"]] == ["XG"]

    def test_a_folder_without_annotation_is_reported(self, tmp_path):
        (tmp_path / "RP003").mkdir()
        (tmp_path / "RP003" / "statement.pdf").write_bytes(b"%PDF")
        p = run_campaign(tmp_path, extractor({})).packages[0]
        assert p.status == "UNREADABLE"

    def test_the_results_folder_is_not_a_package(self, tmp_path):
        package(tmp_path, "RP001", BALANCED)
        (tmp_path / "_results").mkdir()
        campaign = run_campaign(tmp_path, extractor({"RP001": BALANCED}))
        assert [p.package for p in campaign.packages] == ["RP001"]
        assert "RP001/annotation_RP001.xlsx" in campaign.inputs


    # --- exercise 20: what is stored --------------------------------------
    def test_the_stored_campaign_is_plain_json(self, tmp_path):
        import json

        from app.services.corpus_campaign import to_entity

        package(tmp_path, "RP001", BALANCED)
        read = dict(BALANCED, XB=22_000_000)
        entity = to_entity(run_campaign(tmp_path, extractor({"RP001": read})), author="Helena")
        json.dumps(entity.summary)  # no Decimal left: it raises otherwise
        json.dumps(entity.packages)
        assert entity.kind == "CORPUS" and entity.author == "Helena"
        assert entity.summary["amplification"] == "2.80"
        assert entity.packages[0]["silent_false_positives"] == ["XB"]

    # --- the details a failed campaign needs ---------------------------------
    def test_every_item_and_ratio_is_reported_with_its_gap(tmp_path):
        package(tmp_path, "RP001", BALANCED)
        read = dict(BALANCED, XB=22_000_000)  # +10 %
        p = run_campaign(tmp_path, extractor({"RP001": read})).packages[0]

        xb = next(f for f in p.fields if f["code"] == "XB")
        assert (xb["status"], xb["silent"]) == ("WRONG", True)
        assert Decimal(xb["difference"]) == Decimal(2_000_000)
        assert Decimal(xb["relative"]) == Decimal("0.1")
        assert all(f["status"] == "RIGHT" for f in p.fields if f["code"] != "XB")

        r4 = next(r for r in p.ratio_rows if r["code"] == "R4")
        assert r4["right"] is False and r4["caused_by"] == ["XB"]
        assert all(r["right"] for r in p.ratio_rows if r["code"] != "R4")
        assert {c["code"] for c in p.checks} >= {"CHK001", "CHK002"}

    def test_a_rejected_reference_says_which_check_fails(tmp_path):
        package(tmp_path, "RP001", dict(BALANCED, AZ=5_699_223_032))
        p = run_campaign(tmp_path, extractor({"RP001": BALANCED})).packages[0]
        failing = [c for c in p.reference_checks if c["blocking"]]
        assert [c["code"] for c in failing] == ["CHK002"]