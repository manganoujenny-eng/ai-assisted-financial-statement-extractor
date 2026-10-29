"""Excel export of a validated analysis (FR-19).

Same content as the PDF report (``report_content.build_report()``), in the
language of the screen, laid out as a workbook:

1. **Summary** — the case, the key figures, the verdict, strengths,
   weaknesses, what the figures can be relied on for, points of attention;
2. **Ratios** — value (a number Excel can compute with), reading, what it
   measures, interpretation, previous year and change;
3. **Checks** — the five consistency checks;
4. **Validated data** — every validated amount, year N and N-1, as numbers;
5. **AI commentary** — only once an analyst approved it (BR-13).

Sheet names and every text follow ``language`` (en / fr). Amounts are real
numbers with a thousands format, so Excel shows them in the reader's own
number style.
"""

from __future__ import annotations

import math
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .report_content import build_report

NAVY = "002868"
GREY = "54595F"
LINE = Side(style="thin", color="D5DAE1")
FILLS = {
    "header": PatternFill("solid", fgColor=NAVY),
    "band": PatternFill("solid", fgColor="F1F4F9"),
    "good": PatternFill("solid", fgColor="E7F6EC"),
    "bad": PatternFill("solid", fgColor="FDECEC"),
    "neutral": PatternFill("solid", fgColor="F1F2F4"),
    "na": PatternFill("solid", fgColor="FEF6E4"),
}
COLOURS = {"good": "15803D", "bad": "B91C1C", "neutral": GREY, "na": "A16207", "info": NAVY}
WRAP = Alignment(wrap_text=True, vertical="top")
AMOUNT = "#,##0"


def export_analysis(
    case: dict[str, Any],
    dataset: dict[str, Any],
    ratios: list[dict[str, Any]],
    checks: list[dict[str, Any]],
    year_comparison: dict | None = None,
    ai_analysis: dict | None = None,
    labels=lambda code: code,
    language: str = "en",
    reliability: dict | None = None,
    generated_at: datetime | None = None,
) -> BytesIO:
    """The workbook in memory."""
    report = build_report(case, dataset, ratios, checks, year_comparison, ai_analysis,
                          labels, language, reliability, generated_at)
    t = report["t"]
    book = Workbook()
    _summary(book.active, report)
    _ratios(book.create_sheet(t["sheet.ratios"]), report)
    _checks(book.create_sheet(t["sheet.checks"]), report)
    _amounts(book.create_sheet(t["sheet.amounts"]), report)
    if report["ai"]:
        _ai(book.create_sheet(t["sheet.ai"]), report)

    for sheet in book.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


# --- helpers ------------------------------------------------------------------
def _width(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _text_row(sheet, row: int, text: str, columns: int, chars_per_line: int,
              font: Font | None = None, fill: PatternFill | None = None) -> int:
    """A paragraph across ``columns`` merged cells, tall enough to read
    (Excel never sizes a merged row by itself)."""
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = sheet.cell(row=row, column=1, value=text)
    cell.alignment = WRAP
    if font:
        cell.font = font
    if fill:
        for column in range(1, columns + 1):
            sheet.cell(row=row, column=column).fill = fill
    lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\n"))
    sheet.row_dimensions[row].height = max(16, 15 * lines + 3)
    return row + 1


def _heading(sheet, row: int, text: str, columns: int, colour: str = NAVY, size: int = 12) -> int:
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, size=size, color=colour)
    for column in range(1, columns + 1):
        sheet.cell(row=row, column=column).border = Border(bottom=Side(style="medium", color=colour))
    sheet.row_dimensions[row].height = 20
    return row + 1


def _table_header(sheet, row: int, titles: list[str]) -> None:
    for column, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = FILLS["header"]
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 20
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def _border_row(sheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.border = Border(bottom=LINE)
        if cell.alignment is None or not cell.alignment.wrap_text:
            cell.alignment = WRAP


# --- the sheets ---------------------------------------------------------------
def _summary(sheet, r: dict) -> None:
    t, h = r["t"], r["headings"]
    sheet.title = t["sheet.summary"]
    columns, chars = 4, 125
    _width(sheet, [26, 30, 26, 30])

    sheet.cell(row=1, column=1, value=t["brand"]).font = Font(size=8, color=GREY)
    sheet.cell(row=2, column=1, value=r["title"]).font = Font(bold=True, size=18, color=NAVY)
    sheet.row_dimensions[2].height = 26
    sheet.cell(row=3, column=1, value=r["subtitle"]).font = Font(size=11, color=GREY)
    row = 5

    facts = r["facts"]
    for i in range(0, len(facts), 2):
        for offset, (label, value) in enumerate(facts[i:i + 2]):
            sheet.cell(row=row, column=1 + 2 * offset, value=label).font = Font(color=GREY)
            sheet.cell(row=row, column=2 + 2 * offset, value=value).font = Font(bold=True)
        row += 1
    row += 1

    if r["key_figures"]:
        row = _heading(sheet, row, f'{h["keyFigures"]} (FCFA)', columns)
        previous = r["previous_year"]
        titles = ["", t["yearN"].format(year=r["year"])]
        if previous:
            titles.append(t["yearN"].format(year=previous))
        for column, title in enumerate(titles, start=1):
            sheet.cell(row=row, column=column, value=title).font = Font(bold=True, color=GREY)
        row += 1
        amounts = {a["code"]: a for a in r["amounts"]}
        for kf in r["key_figures"]:
            sheet.cell(row=row, column=1, value=kf["label"])
            current = sheet.cell(row=row, column=2, value=_float(amounts[kf["code"]]["current"]))
            current.number_format, current.font = AMOUNT, Font(bold=True, size=12, color=NAVY)
            if previous:
                before = sheet.cell(row=row, column=3, value=_float(amounts[kf["code"]]["previous"]))
                before.number_format, before.font = AMOUNT, Font(color=GREY)
            _border_row(sheet, row, 3 if previous else 2)
            row += 1
        row += 1

    row = _heading(sheet, row, h["summary"], columns)
    row = _text_row(sheet, row, r["verdict"], columns, chars, Font(size=11), FILLS["band"]) + 1
    for key, colour in (("strengths", "good"), ("weaknesses", "bad"),
                        ("reliance", "info"), ("attention", "info")):
        cell = sheet.cell(row=row, column=1, value=h[key])
        cell.font = Font(bold=True, size=11, color=COLOURS[colour])
        row += 1
        for line in r[key] or [t["none"]]:
            row = _text_row(sheet, row, f"•  {line}", columns, chars)
        row += 1

    row = _heading(sheet, row, h["method"], columns, GREY, 10)
    _text_row(sheet, row, r["method"], columns, chars, Font(size=9, color=GREY))


def _ratios(sheet, r: dict) -> None:
    t = r["t"]
    titles = [t["code"], t["ratio"], t["value"], "", t["reading"], t["meaning"],
              t["interpretation"], t["previous"], t["change"]]
    _width(sheet, [7, 30, 11, 5, 15, 42, 60, 15, 11])
    _table_header(sheet, 1, titles)
    for row, ratio in enumerate(r["ratios"], start=2):
        values = [ratio["code"], ratio["name"],
                  float(ratio["number"]) if ratio["number"] is not None else "—",
                  ratio["unit"], ratio["reading_text"], ratio["meaning"], ratio["interpretation"],
                  ratio["previous"], ratio["change"]]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value).alignment = WRAP
        sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet.cell(row=row, column=3).number_format = "0.00"
        reading = sheet.cell(row=row, column=5)
        reading.fill = FILLS[ratio["reading"]]
        reading.font = Font(bold=True, color=COLOURS[ratio["reading"]])
        lines = max(math.ceil(len(ratio["interpretation"]) / 58), math.ceil(len(ratio["meaning"]) / 40), 1)
        sheet.row_dimensions[row].height = 15 * lines + 4
        _border_row(sheet, row, len(titles))


def _checks(sheet, r: dict) -> None:
    t = r["t"]
    titles = [t["code"], t["check"], t["severity"], t["result"], t["detail"]]
    _width(sheet, [10, 38, 15, 12, 70])
    _table_header(sheet, 1, titles)
    for row, check in enumerate(r["checks"], start=2):
        values = [check["code"], check["label"], check["severity"], check["result"], check["message"]]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value).alignment = WRAP
        state = "good" if check["passed"] else "bad"
        result = sheet.cell(row=row, column=4)
        result.fill, result.font = FILLS[state], Font(bold=True, color=COLOURS[state])
        sheet.row_dimensions[row].height = 15 * max(1, math.ceil(len(check["message"]) / 68)) + 4
        _border_row(sheet, row, len(titles))


def _amounts(sheet, r: dict) -> None:
    t = r["t"]
    titles = [t["code"], t["item"], t["yearN"].format(year=r["year"])]
    if r["has_prior"]:
        titles.append(t["yearN"].format(year=r["year"] - 1))
    _width(sheet, [8, 46, 20, 20])
    _table_header(sheet, 1, titles)
    for row, item in enumerate(r["amounts"], start=2):
        sheet.cell(row=row, column=1, value=item["code"]).font = Font(bold=True)
        sheet.cell(row=row, column=2, value=item["label"])
        current = sheet.cell(row=row, column=3, value=_float(item["current"]))
        current.number_format = AMOUNT
        if r["has_prior"]:
            previous = sheet.cell(row=row, column=4, value=_float(item["previous"]))
            previous.number_format = AMOUNT
        _border_row(sheet, row, len(titles))


def _ai(sheet, r: dict) -> None:
    ai, h = r["ai"], r["headings"]
    _width(sheet, [30, 30, 30, 30])
    columns, chars = 4, 115
    sheet.cell(row=1, column=1, value=h["ai"]).font = Font(bold=True, size=16, color=NAVY)
    row = _text_row(sheet, 2, ai["note"], columns, chars, Font(italic=True, color=GREY)) + 1
    for title, text in ai["paragraphs"]:
        sheet.cell(row=row, column=1, value=title).font = Font(bold=True, size=11, color=NAVY)
        row = _text_row(sheet, row + 1, text, columns, chars) + 1
    for title, points in ai["lists"]:
        sheet.cell(row=row, column=1, value=title).font = Font(bold=True, size=11, color=NAVY)
        row += 1
        for point in points:
            row = _text_row(sheet, row, f"•  {point}", columns, chars)
        row += 1


def _float(value):
    return float(value) if value is not None else None