from openpyxl import load_workbook

def load_excel(path):
    # read_only streams rows instead of building every cell object in memory
    # first — the difference between a few MB and a few GB on a large
    # workbook. Same values (checked on the test workbooks). The caller must
    # close() it when done, or the file stays open on Windows.
    workbook = load_workbook(path, data_only=True, read_only=True)
    return workbook
def extract_excel_text(workbook):
    sheets_data = {}
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(row)
        sheets_data[sheet_name] = rows

    return sheets_data

def excel_cell_to_text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (int, float)):
        return str(value)
    return str(value).strip()

def excel_sheet_to_table(rows):
    return [[excel_cell_to_text(cell) for cell in row] for row in rows]